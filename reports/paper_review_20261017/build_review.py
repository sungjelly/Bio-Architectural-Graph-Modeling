"""Build the reviewed manuscript and its self-contained editorial change record.

Run from the BAGM root with the isolated document environment. No model code or
data is modified. Source text is exported as JSON for future revisions.
"""
from pathlib import Path
import json, hashlib, re
from xml.sax.saxutils import escape
import pymupdf as fitz
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Flowable
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
INK, BLUE, MUTED = '#172E40', '#176F79', '#526575'
WIDTH = 503.28
FONTDIR = Path('/usr/share/fonts/truetype/freefont')
for name, fname in [('Body','FreeSerif.ttf'),('BodyBold','FreeSerifBold.ttf'),('Sans','FreeSans.ttf'),('SansBold','FreeSansBold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(FONTDIR/fname)))
pdfmetrics.registerFontFamily('Body',normal='Body',bold='BodyBold',italic='Body',boldItalic='BodyBold')
pdfmetrics.registerFontFamily('Sans',normal='Sans',bold='SansBold',italic='Sans',boldItalic='SansBold')
pdfmetrics.registerFont(TTFont('Korean',str(ROOT/'reports/paper_draft_20261016/assets/NanumGothic-Regular.ttf')))

def p(text): return ('p',text)
def h(text): return ('h',text)
def eq(text): return ('eq',text)

PAGES = [
[
 ('ko','BAGM을 이용한 위 조직의 공간 세포 표현 학습'),
 ('title','BAGM: Learning Spatial Cell Representations in Gastric Tissue'),
 ('meta','Revised draft · 17 October 2026 · Implementation and evidence reviewed through 12:49 UTC'),
 h('Abstract'),
 p('Spatial transcriptomics enables cell states to be studied within their tissue context. We introduce Bio-Architectural Graph Modeling (BAGM), which learns cell representations by predicting randomly masked expression features from visible RNA measurements, morphology/imaging covariates and spatial neighborhoods. Relative geometry contributes an additive bias to scaled dot-product graph attention, and training minimizes Huber loss at masked entries. Exploratory clustering analyses show changes in spatial partitions between initial and final cell representations; reported visual correspondence with putative cancer regions awaits independent annotation and matched baseline evaluation. A completed donor-separated validation experiment improved masked Huber loss and mean absolute error over constant predictors, while mean squared error was slightly worse. A separate, ongoing paper-specific experiment holds out cores while permitting donor overlap. These observations motivate controlled predictive and region-level evaluation, followed by attention and sensitivity analyses of model-derived dependencies.'),
 h('1. Introduction'),
 p('Spatial transcriptomics connects molecular measurements with the organization of intact tissue. Platforms such as the CosMx Spatial Molecular Imager support highly multiplexed RNA and protein profiling with subcellular localization [1]. Here, BAGM uses RNA expression and selected morphology/imaging covariates from gastric tissue; it does not jointly model a protein-expression assay.'),
 p('Self-supervised learning provides representations without exhaustive cell annotations. Masked reconstruction is established in computer vision [2], while scGPT uses generative pretraining for cellular expression data [3]. Spatial methods incorporate tissue organization: GraphST combines graph-based contrastive learning with expression reconstruction [4], and SpaMask [5] and the SAGE-FM preprint [6] use masked prediction in spatial representations. These studies motivate learning from both molecular measurements and neighborhood structure.'),
 p('An informative clustering representation does not itself establish molecular interactions. Predicting hidden expression supplies a complementary task for examining conditional predictive dependencies. MISTy illustrates how information from different spatial contexts can be used to investigate molecular associations [7]. BAGM similarly aims to support tissue description and subsequent analysis of candidate dependencies, with biological interpretation requiring additional evidence.'),
 p('We examine whether BAGM representations delineate putative tumor-associated tissue domains and how those representations change across graph layers. A node represents a cell, an edge a candidate spatial relationship, and an edge feature measured relative geometry. These choices provide a biological inductive bias; they do not establish that every edge is a direct cellular interaction. The present manuscript separates implemented methods, completed exploratory observations and evaluation that remains pending.'),
],
[
 h('2. Methods'),
 p('<b>Data and covariates.</b> The analyzed cohort contains 407,659 mapped cells from 28 gastric tissue cores across two CosMx slides and 14 donor groups. Inputs comprise 1,000 biological RNA genes and 22 morphology/imaging features. Clinical labels, cell identifiers, slide identifiers, absolute coordinates and full-count library sizes are not node covariates. Morphology preprocessing is fitted on training cells and then frozen.'),
 p('<b>Current paper-specific split.</b> The fresh seed-0 experiment uses the whole-core partition below. Donor overlap is permitted, both test cores are from the second slide, and these samples were examined in earlier experiments. This is a holdout for the new fit, not new-donor validation or a historically untouched confirmatory test. Preprocessing is refitted on the 21 training cores only; validation selects the checkpoint, and test evaluation follows checkpoint selection.'),
 ('split',None),
 p('<b>Expression transformation and masking.</b> We use standardized quasi-Poisson Pearson residuals, related to the residual-normalization framework in [8], with dispersion factor 1.01 and an upper clip of 10 after standardization. This is an expression transformation, not evidence of batch-effect removal. Each cell independently draws a number of hidden genes uniformly from 0 to 1,000 inclusive; positions are sampled without replacement. There are ten training views per core per epoch and three fixed evaluation views per core. Morphology remains visible. Each input view is recomputed from visible raw counts and frozen training statistics so hidden counts cannot influence input normalization. Fully observed transformed features are supervised targets only. The task predicts masked residual features, not raw transcript counts.'),
 p('<b>Spatial graph and model.</b> Coordinate-based graphs are constructed within each core by deterministic neighbor selection within radial distance bands up to 500 μm, with a nominal 200-neighbor budget before bidirectional union. There are no self-loops or cross-core edges. Each edge carries 70 relative-geometry features. These graphs include broad neighborhood context and must not be equated with direct contact. A shared cell encoder combines visible expression, mask indicators and morphology. Four independently parameterized graph Transformer blocks [9, 10] update the cell states, and a decoder predicts expression features (Fig. 1; Appendix A).'),
 ('figure',None),
 ('caption','Figure 1. Current BAGM architecture. Relative geometry adds a bias to attention logits. Values are projected cell states, including morphology and mask information and, in later blocks, accumulated neighborhood context. Residual paths retain same-cell information. Random masking is sampled for every cell; the sampled mask can be empty. The decoder predicts residual features, with loss evaluated only at masked entries.'),
],
[
 h('2. Methods (continued)'),
 p('<b>Objective and evaluation.</b> Huber loss with δ = 1 is averaged over masked entries within each view, then equally over cores and views. AdamW uses learning rate 10<super>−4</super>, weight decay 10<super>−5</super> and gradient-norm clipping at 1. The maximum is 300 epochs; early stopping uses equal-core validation masked Huber loss, a minimum of 50 epochs and patience of 25 validations without improvement exceeding 10<super>−4</super>. The best checkpoint tracks strict validation-loss improvements. Test masks use a separate fixed namespace. Cells are correlated observations; donors are biological replication units, and cores are the current holdout units.'),
 p('<b>Representation analysis.</b> Leiden clustering [11] is applied to learned representations, including the initial state h<super>(0)</super> and final state h<super>(4)</super>. Fully observed embedding extraction is a separate analysis from masked prediction. PCA-based representations provide a descriptive comparator. To claim improved cancer-region delineation, clustering settings must be matched and maps scored against independently recorded region annotations.'),
 h('3. Results and current status'),
 p('<b>Exploratory tissue representations.</b> Completed analyses show that the initial and final representations yield different spatial partitions. For the completed joint-slide model, fully observed embeddings from all 407,659 cells produced 16 initial-state and 22 final-state primary clusters, with adjusted Rand index 0.370 between the partitions. This joint analysis includes training and validation cells. It establishes a change in representation, not tumor-boundary accuracy or independent biological replication.'),
 p('Preliminary visual inspection by the authors suggested clearer cancer-region correspondence after four layers and relative to PCA. These observations remain qualitative. Independent clinician region masks and a matched quantitative comparison were not available in the documented analyses. Existing PCA and embedding clustering pipelines also differ in settings. Consequently, improved cancer delineation is not established, and differences between h<super>(0)</super> and h<super>(4)</super> cannot be attributed specifically to geometry rather than attention, feed-forward transformations, residual updates or smoothing.'),
 p('<b>Completed predecessor prediction experiment.</b> A separate single-seed experiment used 20 training cores from ten donor groups and eight validation cores from four other donor groups, with no independent test partition. Validation cores had already been examined during preprocessing diagnostics, so the results remain exploratory. The validation-selected model achieved equal-core masked Huber loss 0.196114, pooled mean absolute error 0.308160, pooled mean squared error 1.005888 and pooled Pearson correlation 0.134499. Across three constant references, Huber loss improved by 6.21–6.72% and mean absolute error by 26.8–28.2%, while mean squared error was 0.58–0.70% worse. These are mixed exploratory validation results; they do not establish graph-specific benefit. They belong to the completed 20/8-core study, not the ongoing 21/5/2-core fit.'),
 p('<b>Pending paper-specific evaluation.</b> At the review cutoff, the current 21/5/2-core experiment was still training; its final test evaluation and test-only Leiden analysis were pending. Current Pearson-feature regression is evaluated using regression metrics. Earlier nonzero-count detection diagnostics used other models and preprocessing; their precision, recall or average precision must not be presented as results of this run. No quantitative cancer-region accuracy is reported here.'),
],
[
 h('4. Discussion and conclusion'),
 p('BAGM provides an exploratory framework for learning spatial cell representations and predicting masked expression features. The available evidence supports changes in representation and limited predictive improvement over constants. Partial-gene prediction may exploit within-cell co-expression, so this gain alone does not establish a benefit from neighborhood information. It does not yet establish improved cancer-region delineation, a graph-specific predictive gain or a cellular mechanism. Larger datasets may improve precision of evaluation, but they do not replace appropriate controls or independent biological replication.'),
 p('Priorities are matched cell-autonomous, morphology, neighborhood-composition and spatial-smoothing baselines; graph and geometry ablations; multiple training seeds; and independent region annotations. Region-level Dice/IoU and boundary agreement require a predefined mapping from clusters to regions. Patient-held-out evaluation and uncertainty across biological units are needed for new-patient generalization. Cell composition, broad tissue fields, core/slide effects and segmentation spillover remain alternative explanations. Binary cross-entropy or hurdle formulations would address different targets and require controlled comparisons; the present evidence does not identify Huber loss as the cause of a detection deficit.'),
 p('Attention describes normalized computational routing and is not, by itself, biological importance. Planned attribution and Jacobian analyses will quantify sensitivity of selected predicted features to visible transformed input features, with morphology, masks and graphs held fixed. Raw-count derivatives are different because they include the exposure-normalization pathway. One-hop effects must be distinguished from multihop dependencies. Bounded input perturbations and message ablations assess faithfulness; separate tests assess stability and calibration against realistic nulls. Independent molecular evidence and controlled perturbation are needed to support mechanism and causality, respectively. Biological interpretation should focus on targets with reproducible graph-specific predictive gain.'),
 h('Appendix A. Current graph attention architecture'),
 p('<b>Cell encoder.</b> Separate linear projections of the visible expression vector, mask and 22 morphology/imaging covariates are summed with a learned bias to form a 256-dimensional state, followed by LayerNorm, GELU and dropout. Hidden expression entries are zero. Four independently parameterized pre-normalized graph blocks each use eight heads of dimension 32.'),
 p('<b>Spatial attention.</b> For sender j, receiver i and head a, query q, key k and value v are bias-free projections of LayerNorm-transformed cell states. Relative geometry ρ<sub>ij</sub> contains 64 distance-basis values plus six distance/orientation/anisotropy descriptors. A 70 → 128 → LayerNorm → GELU → 8 network produces one additive bias per head:'),
 eq('s<sub>ij</sub><super>a</super> = (q<sub>i</sub><super>a</super>)<super>T</super>k<sub>j</sub><super>a</super> / √32 + b<sub>a</sub>(ρ<sub>ij</sub>)'),
 eq('α<sub>ij</sub><super>a</super> = softmax<sub>j∈N(i)</sub>(s<sub>ij</sub><super>a</super>),     z<sub>i</sub><super>a</super> = Σ<sub>j∈N(i)</sub> α<sub>ij</sub><super>a</super> v<sub>j</sub><super>a</super>'),
 p('The geometry output weights and biases start at zero. The current model uses queries and keys without additional L2 normalization, with fixed 1/√32 scaling; it has no geometry feature gate or learned attention temperature. Geometry directly enters logits only. Queries, keys and values are recomputed at each layer. The incoming-neighbor softmax has no self-edge coefficient; same-cell information is retained by residual paths.'),
 p('<b>Updates and decoding.</b> Concatenated heads pass through a bias-free 256-dimensional output projection, dropout and residual addition. A second pre-normalized 256 → 1,024 → 256 GELU feed-forward transformation has another residual connection. The 256 → 1,024 → 1,000 decoder uses GELU and dropout, with unconstrained outputs. Dropout is 0.1 and attention dropout is zero. For transformed target Y and mask M:'),
 eq('L<sub>view</sub> = [Σ<sub>ig</sub> M<sub>ig</sub> Huber<sub>1</sub>(Ŷ<sub>ig</sub> − Y<sub>ig</sub>)] / [Σ<sub>ig</sub> M<sub>ig</sub>]'),
],
[
 h('References'),
 ('ref','[1] S. He, R. Bhatt, C. Brown, et al. “High-plex imaging of RNA and proteins at subcellular resolution in fixed tissue by spatial molecular imaging.” <i>Nature Biotechnology</i> 40, 1794–1806 (2022). <link href="https://doi.org/10.1038/s41587-022-01483-z" color="#176F79">doi:10.1038/s41587-022-01483-z</link>'),
 ('ref','[2] K. He, X. Chen, S. Xie, et al. “Masked Autoencoders Are Scalable Vision Learners.” <i>CVPR</i>, 16000–16009 (2022). <link href="https://openaccess.thecvf.com/content/CVPR2022/html/He_Masked_Autoencoders_Are_Scalable_Vision_Learners_CVPR_2022_paper.html" color="#176F79">Official proceedings</link>.'),
 ('ref','[3] H. Cui, C. Wang, H. Maan, et al. “scGPT: toward building a foundation model for single-cell multi-omics using generative AI.” <i>Nature Methods</i> 21, 1470–1480 (2024). <link href="https://doi.org/10.1038/s41592-024-02201-0" color="#176F79">doi:10.1038/s41592-024-02201-0</link>'),
 ('ref','[4] Y. Long, K. S. Ang, M. Li, et al. “Spatially informed clustering, integration, and deconvolution of spatial transcriptomics with GraphST.” <i>Nature Communications</i> 14, 1155 (2023). <link href="https://doi.org/10.1038/s41467-023-36796-3" color="#176F79">doi:10.1038/s41467-023-36796-3</link>'),
 ('ref','[5] W. Min, D. Fang, J. Chen and S. Zhang. “SpaMask: Dual masking graph autoencoder with contrastive learning for spatial transcriptomics.” <i>PLOS Computational Biology</i> 21(4), e1012881 (2025). <link href="https://doi.org/10.1371/journal.pcbi.1012881" color="#176F79">doi:10.1371/journal.pcbi.1012881</link>'),
 ('ref','[6] X. Zhan, J. Xu, Y. Zheng, Z. Good and O. Gevaert. “SAGE-FM: A lightweight and interpretable spatial transcriptomics foundation model.” <i>arXiv preprint</i> arXiv:2601.15504 (2026). <link href="https://arxiv.org/abs/2601.15504" color="#176F79">arxiv.org/abs/2601.15504</link>'),
 ('ref','[7] J. Tanevski, R. O. Ramirez Flores, A. Gabor, et al. “Explainable multiview framework for dissecting spatial relationships from highly multiplexed data.” <i>Genome Biology</i> 23, 97 (2022). <link href="https://doi.org/10.1186/s13059-022-02663-5" color="#176F79">doi:10.1186/s13059-022-02663-5</link>'),
 ('ref','[8] J. Lause, P. Berens and D. Kobak. “Analytic Pearson residuals for normalization of single-cell RNA-seq UMI data.” <i>Genome Biology</i> 22, 258 (2021). <link href="https://doi.org/10.1186/s13059-021-02451-7" color="#176F79">doi:10.1186/s13059-021-02451-7</link>'),
 ('ref','[9] P. Veličković, G. Cucurull, A. Casanova, et al. “Graph Attention Networks.” <i>ICLR</i> (2018). <link href="https://arxiv.org/abs/1710.10903" color="#176F79">arxiv.org/abs/1710.10903</link>'),
 ('ref','[10] A. Vaswani, N. Shazeer, N. Parmar, et al. “Attention Is All You Need.” <i>Advances in Neural Information Processing Systems</i> 30 (2017). <link href="https://papers.nips.cc/paper/7181-attention-is-all-you-need" color="#176F79">Official proceedings</link>.'),
 ('ref','[11] V. A. Traag, L. Waltman and N. J. van Eck. “From Louvain to Leiden: guaranteeing well-connected communities.” <i>Scientific Reports</i> 9, 5233 (2019). <link href="https://doi.org/10.1038/s41598-019-41695-z" color="#176F79">doi:10.1038/s41598-019-41695-z</link>'),
 ('note','References [9, 10] provide the graph-attention and Transformer background for the architecture in Appendix A. The revised manuscript retains all eleven references from the supplied paper. The following pages are an editorial change record, separate from the manuscript.'),
],
[
 ('meta','EDITORIAL CHANGE RECORD · 17 OCTOBER 2026'),
 h('Review outcome'),
 p('The supplied paper agrees with the broad research aim, but not fully with the current implementation or evidence. The revised manuscript describes the current original Relative-QKV architecture, separates three distinct experiment contexts, and calibrates biological claims. This was a review of code, protocols and existing aggregate results; no model was retrained or new scientific result generated.'),
 h('Material wording and technical changes'),
 p('<b>1. Geometry “feature modulation and bias” (Fig. 1; Appendix A).</b> Replaced with additive geometry bias in scaled dot-product attention. Removed the feature gates, L2-normalized queries/keys, learned temperature and tanh biases. Those details describe the historical geometry-modulated variant, not the current paper model. Replaced the equation and redrew the schematic to match the active implementation. [S1]'),
 p('<b>2. “Metadata” and “molecular content exchanged” (Methods; Fig. 1).</b> Specified the 22 permitted morphology/imaging covariates. Clarified that value vectors contain encoded RNA, mask and morphology information, with accumulated graph context in later blocks. Clinical labels and identifiers are not model inputs. [S1, S2]'),
 p('<b>3. “Normalized using Pearson residuals … across slides” (Methods).</b> Named the implemented quasi-Poisson transform, train-only fitting and visible-count input normalization. Clarified that targets are standardized residual features and normalization does not demonstrate batch correction. The cited residual paper provides conceptual background, not evidence that this exact preprocessing removes slide effects. [S2; ref. 8]'),
 p('<b>4. Missing split and experiment status (Methods; Results).</b> Added the current 21/5/2-core split and cell counts, donor overlap, previous sample use and review cutoff. Distinguished it from the completed 20/8-core donor-separated validation experiment. The current final test and test-only clustering remain pending. [S2, S3]'),
 p('<b>5. “Clearer … correspondence with clinician-identified cancer regions than PCA” (Abstract; Results).</b> Changed to preliminary author-reported visual impressions. Independent region masks and a matched quantitative PCA comparison were not available in the reviewed records. No superiority, pathology validation or segmentation accuracy is asserted. [S2, S4]'),
 p('<b>6. “Improved delineation after four attention layers” (Abstract; Results).</b> Replaced with a verified change in clustering partitions. Explained that h0/h4 comparisons do not isolate geometry and that the completed joint embedding analysis includes training and validation cells. [S4]'),
 p('<b>7. “Precision, recall, and AUPRC … remained limited under Huber-loss training” (Abstract; Results).</b> Removed this unscoped attribution. Historical detection diagnostics used different models, input scales and evaluation definitions; no matching AUPRC result for the current Pearson regression run was verified. Added the completed predecessor’s Huber, MAE, MSE and Pearson results, including worse MSE, without assigning them to the active paper fit. [S3, S5]'),
],
[
 ('meta','EDITORIAL CHANGE RECORD · CONTINUED'),
 p('<b>8. “RNA expression” and “reconstructs RNA measurements” (throughout).</b> Used “expression features” or “standardized residual features” where that is the actual prediction target. Preserved RNA as the source measurement. This avoids implying raw-count recovery or a count-likelihood model. [S2]'),
 p('<b>9. “All cells masked” (Fig. 1).</b> Replaced with masking sampled independently for every cell. The number of hidden genes is uniform from 0 through 1,000; an empty mask and a whole-cell mask are possible. This is not a dedicated whole-node benchmark. [S2]'),
 p('<b>10. Attention, Jacobians and “dependencies” (Discussion).</b> Specified computational routing and local model sensitivity, the input scale and fixed mask/graph/morphology. Added one-hop versus multihop distinctions and the need for matched predictive controls, faithfulness, nulls and independent/perturbational evidence before mechanism or causal claims. [S1, S2]'),
 p('<b>11. Evaluation priorities (Discussion).</b> Added matched self/context baselines, graph/geometry ablations, multiple training seeds and independent region scoring. Preserved the caution that Huber loss has not been shown to cause the detection limitation; BCE/hurdle models change the target and require controlled comparisons. [S2, S3, S5]'),
 h('Evidence and traceability'),
 p('Source is the supplied three-page BAGM_Paper.pdf. The review used the current working tree and existing aggregate reports; it did not rerun the experiments or independently score histology. Completed performance and representation results are attributed to the predecessor experiment, and the current-run status is explicitly time-bounded. No patient-level records are included.'),
 ('source','S1 — Current architecture: src/spatial_benchmark/relative_qkv_graph_transformer.py (attention formula and geometry encoder); src/spatial_benchmark/models.py (cell encoder and decoder). Historical gated variant: experiments/campaigns/cmp_20260903_so2_14core_geometry_modulated_relative_qkv_seed0_batch2/README.md.'),
 ('source','S2 — Current paper protocol: experiments/campaigns/cmp_20261017_so12_paper_core_split/README.md and configs/seed0.yaml. Normalization: src/spatial_benchmark/so12_relative_qkv_data.py. Active run at review: r_20261017T044837Z_caf76350_s000_f00_a01_12df41be.'),
 ('source','S3 — Completed validation: experiments/campaigns/cmp_20261016_so12_pearson_relative_qkv_train_validation/README.md, final corrected-checkpoint results. Immutable run: r_20261016T090820Z_ad267085_s000_f00_a02_451ef28b; selected epoch 151. Three constant references: zero, equal-core training mean and equal-cell training mean.'),
 ('source','S4 — Completed representation analysis: experiments/campaigns/cmp_20261017_so12_h0_hl_leiden_segmentation/README.md. Pending test-only work: experiments/campaigns/cmp_20261017_so12_paper_test_leiden/README.md. These do not establish independent pathology agreement.'),
 ('source','S5 — Historical detection: experiments/campaigns/cmp_20260905_so2_nonzero_accuracy/README.md and cmp_20260906_so2_failure_diagnosis/README.md. Fitted-cohort standardized-log1p count detection is distinct from the current held-out residual-feature regression.'),
 ('note','Editorial scope: the core research question, bilingual title, section structure and all eleven references are retained. The layout is expanded for legibility, equations are rebuilt, and a change record is appended. Remaining research requirements are identified explicitly rather than represented as completed findings.'),
],
]

class Architecture(Flowable):
    def __init__(self):
        super().__init__()
        self.width=WIDTH
        self.height=104
    def draw(self):
        c=self.canv
        c.setFillColor(colors.HexColor('#F3F7F9')); c.roundRect(0,0,self.width,self.height,5,fill=1,stroke=0)
        def box(x,y,w,h,lines,color):
            c.setFillColor(colors.HexColor(color)); c.roundRect(x,y,w,h,4,fill=1,stroke=0)
            c.setFillColor(colors.white);c.setFont('SansBold',8)
            for j,line in enumerate(lines):c.drawCentredString(x+w/2,y+h/2+(len(lines)-1)*5-j*10-2,line)
        def arrow(x,y,x2,y2):
            c.setStrokeColor(colors.HexColor(BLUE));c.setLineWidth(.8);c.line(x,y,x2,y2)
            c.line(x2,y2,x2-4,y2+2);c.line(x2,y2,x2-4,y2-2)
        box(9,37,88,34,['Visible RNA + mask','Morphology'],'#476B88')
        box(112,37,65,34,['Cell encoder','h0 · 256-D'],'#2674AE')
        box(194,37,137,34,['4 graph Transformer blocks','dot product + geometry bias'],'#2674AE')
        box(348,37,65,34,['Decoder','h4 → features'],'#148C88')
        box(430,37,64,34,['Masked-only','Huber loss'],'#A16814')
        for a,b in [(97,112),(177,194),(331,348),(413,430)]:arrow(a,54,b,54)
        c.setFillColor(colors.HexColor(INK));c.setFont('Sans',8)
        c.drawCentredString(262,87,'Within-core spatial graph · no self-loops')
        c.drawCentredString(262,18,'70-D relative geometry → 8 additive attention biases')
        c.setStrokeColor(colors.HexColor(BLUE));c.line(262,28,262,37)
        c.line(262,82,262,71)

styles={
 'p':ParagraphStyle('p',fontName='Body',fontSize=10.4,leading=13.35,textColor=colors.HexColor(INK),spaceAfter=7),
 'h':ParagraphStyle('h',fontName='SansBold',fontSize=12.5,leading=16,textColor=colors.HexColor(BLUE),spaceBefore=7,spaceAfter=7,keepWithNext=True),
 'title':ParagraphStyle('title',fontName='SansBold',fontSize=20,leading=24,textColor=colors.HexColor(INK),spaceAfter=10),
 'ko':ParagraphStyle('ko',fontName='Korean',fontSize=12.5,leading=17,textColor=colors.HexColor(INK),spaceAfter=8),
 'meta':ParagraphStyle('meta',fontName='Sans',fontSize=8,leading=11,textColor=colors.HexColor(MUTED),spaceAfter=10),
 'caption':ParagraphStyle('caption',fontName='Body',fontSize=9,leading=11.4,textColor=colors.HexColor(MUTED),spaceBefore=6,spaceAfter=6),
 'eq':ParagraphStyle('eq',fontName='Body',fontSize=11.7,leading=18,alignment=1,textColor=colors.HexColor(INK),spaceBefore=2,spaceAfter=6),
 'ref':ParagraphStyle('ref',fontName='Body',fontSize=10.4,leading=13.5,textColor=colors.HexColor(INK),spaceAfter=12),
 'source':ParagraphStyle('source',fontName='Sans',fontSize=8.2,leading=10.9,textColor=colors.HexColor(INK),spaceAfter=8,splitLongWords=True),
 'note':ParagraphStyle('note',fontName='Body',fontSize=9.5,leading=12.4,textColor=colors.HexColor(MUTED),spaceBefore=9,spaceAfter=7),
}

def footer(c,d):
    c.saveState();c.setStrokeColor(colors.HexColor('#CBD7DE'));c.setLineWidth(.5);c.line(46,38,549.28,38)
    c.setFont('Sans',7.5);c.setFillColor(colors.HexColor(MUTED))
    c.drawString(46,25,'BAGM · Revised manuscript' if d.page<=5 else 'BAGM · Editorial change record')
    c.drawRightString(549.28,25,f'{d.page} / 7');c.restoreState()

def split_table():
    rows=[['Partition','Cores','Cells'],['Training','21','264,491'],['Validation','5','61,561'],['Test','2','81,607']]
    tab=Table(rows,colWidths=[WIDTH*.5,WIDTH*.2,WIDTH*.3],hAlign='LEFT')
    tab.setStyle(TableStyle([('FONTNAME',(0,0),(-1,0),'SansBold'),('FONTNAME',(0,1),(-1,-1),'Sans'),('FONTSIZE',(0,0),(-1,-1),9),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E6F0F3')),('TEXTCOLOR',(0,0),(-1,-1),colors.HexColor(INK)),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#9BB2BD')),('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#CBD7DE')),('BOTTOMPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5)]))
    return tab

story=[]
for n,page in enumerate(PAGES):
    if n:story.append(PageBreak())
    for kind,text in page:
        if kind=='figure':story.extend([Spacer(1,4),Architecture()])
        elif kind=='split':story.extend([split_table(),Spacer(1,9)])
        else:story.append(Paragraph(text,styles[kind]))

pdf=OUT/'BAGM_Paper_Revised_with_Changes_2026-10-17.pdf'
doc=SimpleDocTemplate(str(pdf),pagesize=(595.28,841.89),leftMargin=46,rightMargin=46,topMargin=42,bottomMargin=51,
    title='BAGM: Learning Spatial Cell Representations in Gastric Tissue — Revised with Change Record',author='BAGM',subject='Reviewed manuscript and editorial changes, 17 October 2026')
doc.build(story,onFirstPage=footer,onLaterPages=footer)
(OUT/'manuscript_content.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2))

# Editable local source for future author revisions. Google Drive PDF is primary.
wd=Document();sec=wd.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
wd.styles['Normal'].font.name='Times New Roman';wd.styles['Normal'].font.size=Pt(11)
def plain(s):
    return re.sub('<[^>]+>','',s).replace('&gt;','>').replace('&lt;','<').replace('&amp;','&')
for n,page in enumerate(PAGES):
    if n:wd.add_page_break()
    for kind,text in page:
        if kind=='figure':wd.add_paragraph('Visible RNA + mask + morphology → cell encoder → 4 graph blocks (additive geometry bias) → decoder → masked-only Huber loss')
        elif kind=='split':
            t=wd.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
            for c,v in zip(t.rows[0].cells,['Partition','Cores','Cells']):c.text=v
            for row in [['Training','21','264,491'],['Validation','5','61,561'],['Test','2','81,607']]:
                for c,v in zip(t.add_row().cells,row):c.text=v
        elif kind in ['h','title']:wd.add_heading(plain(text),level=0 if kind=='title' else 1)
        else:wd.add_paragraph(plain(text))
wd.save(OUT/'BAGM_Paper_Revised_with_Changes_2026-10-17.docx')

d=fitz.open(pdf);preview=OUT/'previews';preview.mkdir(exist_ok=True)
checks={'pages':len(d),'expected_pages':7,'bytes':pdf.stat().st_size,'page_checks':[],'links':sum(len(pg.get_links()) for pg in d)}
texts=[]
for i,pg in enumerate(d):
    txt=pg.get_text();texts.append(txt)
    outside=[]
    for b in pg.get_text('dict')['blocks']:
        for line in b.get('lines',[]):
            for s in line['spans']:
                box=fitz.Rect(s['bbox'])
                if box.x0 < 40 or box.x1 > 555 or box.y0 < 30 or box.y1 > 825:outside.append({'text':s['text'],'bbox':list(box)})
    checks['page_checks'].append({'page':i+1,'characters':len(txt),'outside_safe_bounds':outside,'replacement_glyphs':txt.count('\ufffd')})
    pg.get_pixmap(matrix=fitz.Matrix(1.35,1.35)).save(preview/f'page_{i+1}.png')
checks['sha256']=hashlib.sha256(pdf.read_bytes()).hexdigest()
checks['page_count_ok']=len(d)==7
checks['layout_ok']=all(not pg['outside_safe_bounds'] and not pg['replacement_glyphs'] for pg in checks['page_checks'])
(OUT/'verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
(OUT/'revised_text.txt').write_text('\n\n'.join(texts))
(OUT/'SHA256SUMS.txt').write_text('\n'.join(f'{hashlib.sha256((OUT/f).read_bytes()).hexdigest()}  {f}' for f in [pdf.name,'BAGM_Paper_Revised_with_Changes_2026-10-17.docx','manuscript_content.json'])+'\n')
print(json.dumps(checks,ensure_ascii=False,indent=2))
