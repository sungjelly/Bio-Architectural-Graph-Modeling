"""Build editable BAGM slides and a compact manuscript diagram (CPU only).

All PowerPoint content is native shapes, connectors, and text. Source model:
src/spatial_benchmark/relative_qkv_graph_transformer.py, original Relative-QKV.
This script reads no biological data and does not access training processes.
"""
from __future__ import annotations

from pathlib import Path
import math
import zipfile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pymupdf

OUT = Path(__file__).resolve().parent
ASSETS = OUT / "assets"
WIDTH, HEIGHT = 13.333333, 7.5
NAVY = "16354A"
BLUE = "247CB5"
TEAL = "218F8D"
PALE = "EAF3F8"
MINT = "E8F5F2"
GRAY = "526470"
RULE = "B9CCD5"
WHITE = "FFFFFF"
LIGHT = "F5F8FA"


def register_fonts():
    root = Path("/usr/share/fonts/truetype/freefont")
    for name, file in [("FreeSans", "FreeSans.ttf"), ("FreeSansBold", "FreeSansBold.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(root / file)))
    for name, file in [("NanumGothic", "NanumGothic-Regular.ttf"),
                       ("NanumGothicBold", "NanumGothic-Bold.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(ASSETS / file)))


def rgb(c):
    return tuple(int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))


class Slide:
    """Draw the same vector layout to a PowerPoint slide and a PDF proof."""
    def __init__(self, prs, pdf, korean=False):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.pdf = pdf
        self.korean = korean
        self.pfont = "NanumGothic" if korean else "FreeSans"

    def rect(self, x, y, w, h, fill=WHITE, stroke=RULE, radius=True, dash=False):
        shp = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shp.adjustments[0] = 0.08
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        shp.line.color.rgb = RGBColor.from_string(stroke); shp.line.width = Pt(1)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        self.pdf.setFillColorRGB(*rgb(fill)); self.pdf.setStrokeColorRGB(*rgb(stroke))
        self.pdf.setLineWidth(1)
        self.pdf.setDash(4, 3) if dash else self.pdf.setDash()
        self.pdf.roundRect(x * 72, (HEIGHT-y-h) * 72, w * 72, h * 72,
                           0.08 * 72 if radius else 0, fill=1, stroke=1)
        self.pdf.setDash()
        return shp

    def text(self, x, y, w, h, text, size=18, color=NAVY, bold=False,
             align="left", font=None, leading=1.18):
        tx = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tx.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = 0; tf.margin_right = 0
        tf.margin_top = 0; tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        use_font = font or self.pfont
        for n, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = use_font; p.font.size = Pt(size)
            p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(0); p.space_before = Pt(0)
            p.line_spacing = Pt(size * leading)
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
        pdf_font = use_font + "Bold" if bold else use_font
        if pdf_font not in pdfmetrics.getRegisteredFontNames():
            pdf_font = "FreeSansBold" if bold else "FreeSans"
        for line in text.split("\n"):
            assert pdfmetrics.stringWidth(line,pdf_font,size) <= w*72+1, f"Text exceeds box: {line}"
        self.pdf.setFont(pdf_font, size); self.pdf.setFillColorRGB(*rgb(color))
        for n, line in enumerate(text.split("\n")):
            xpos = x * 72
            ypos = (HEIGHT-y) * 72-size*0.85-n*size*leading
            if align == "center": self.pdf.drawCentredString(xpos+w*36, ypos, line)
            elif align == "right": self.pdf.drawRightString(xpos+w*72, ypos, line)
            else: self.pdf.drawString(xpos, ypos, line)
        return tx

    def line(self, x1, y1, x2, y2, color=BLUE, width=1.6, arrow=True, dash=False):
        line = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = RGBColor.from_string(color); line.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            end = OxmlElement("a:tailEnd"); end.set("type", "triangle")
            end.set("w", "sm"); end.set("len", "sm")
            line.line._get_or_add_ln().append(end)
        c = self.pdf; c.setStrokeColorRGB(*rgb(color)); c.setFillColorRGB(*rgb(color))
        c.setLineWidth(width); c.setDash(4,3) if dash else c.setDash()
        ax, ay, bx, by = x1*72, (HEIGHT-y1)*72, x2*72, (HEIGHT-y2)*72
        c.line(ax,ay,bx,by); c.setDash()
        if arrow:
            a=math.atan2(by-ay,bx-ax); length=6.5; spread=2.8
            p=c.beginPath(); p.moveTo(bx,by)
            p.lineTo(bx-length*math.cos(a)+spread*math.sin(a),by-length*math.sin(a)-spread*math.cos(a))
            p.lineTo(bx-length*math.cos(a)-spread*math.sin(a),by-length*math.sin(a)+spread*math.cos(a))
            p.close(); c.drawPath(p,fill=1,stroke=0)
        return line

    def notes(self, text):
        self.slide.notes_slide.notes_text_frame.text = text

    def end(self):
        self.pdf.showPage()


def header(s, title, subtitle, page):
    s.text(.5,.32,12.3,.55,title,size=29,bold=True)
    s.text(.52,.96,12.15,.4,subtitle,size=14,color=GRAY)
    s.line(.52,1.43,12.82,1.43,color=RULE,width=1,arrow=False)
    s.text(.52,7.13,11.6,.2,"BAGM  •  Editable architecture draft  •  16 October 2026",size=9,color=GRAY)
    s.text(12.35,7.13,.45,.2,str(page),size=9,color=GRAY,align="right")


NOTES = """Architecture scope: original Relative-Geometric QKV, the active joint SO_1 + SO_2 model chosen by the user. It is NOT the September geometry-modulated cosine-QKV variant. The original model has additive geometry bias only: it does not L2-normalize Q/K, learn a bounded cosine scale, or dimensionwise-modulate Q-K products. Each of four graph blocks has independent parameters. Geometry enters the attention scores only; value projections act on normalized node states. Morphology may affect values indirectly through node states.
Sources: src/spatial_benchmark/relative_qkv_graph_transformer.py; src/spatial_benchmark/models.py; experiments/campaigns/cmp_20261016_so12_pearson_relative_qkv_train_validation/README.md and configs/seed0_attempt2.yaml. Directed graph edges are sender j to receiver i; there are no self-loop edges. Same-cell information travels through residuals and the encoder. The k200/r500 um graph is radial-stratified with symmetrization, so final incoming degree is not necessarily 200.
Current objective: each cell independently masks 0 through 1,000 genes uniformly, then samples that many gene positions without replacement. Input Pearson features are recomputed from visible raw counts only, using training-fitted gene frequencies and transformations. Hidden entries are exactly zero; an explicit gene mask and 22 permitted morphology/imaging fields remain visible. Ten training mask views per core per epoch; three fixed validation views. Fully observed Pearson targets are used only for supervision. Huber loss delta=1 is evaluated at masked entries.
Dataset split: 20 training cores/10 donor groups and 8 validation cores/4 donor groups. Validation is exploratory and used for model selection; there is no untouched test partition. Training status and metrics are intentionally omitted from the architecture diagram. Leiden clustering is an exploratory downstream representation analysis, not part of the masked prediction objective. Attention/Jacobian analyses are future hypothesis generation; attention is model routing and Jacobians are local sensitivities, not causal effects. Historical clustering findings must not be assigned to this active run without new evidence.
All visible slide objects are native editable shapes, connectors, or text boxes. The accompanying architecture_preview.pdf is a vector layout proof generated from the same drawing calls, not a rendering by Microsoft PowerPoint. Fonts: FreeSans and NanumGothic; PowerPoint can substitute another sans-serif font if necessary."""


def overview(prs, pdf, korean=False):
    s=Slide(prs,pdf,korean)
    if korean:
        header(s,"BAGM: 조직 그래프에서 마스킹된 발현 예측",
               "현재 모델: 원형 Relative-Geometric QKV  |  SO_1 + SO_2 통합 데이터",3)
        titles=["마스킹된 세포 입력","노드 인코더","그래프 블록 × 4","발현 디코더"]
        bodies=["관측 유전자만으로 Pearson 계산\n명시적 마스크 + 형태 정보 22개",
                "공유 투영 + 정규화\nh0: 256차원", "서로 독립된 4개 층\n8개 헤드 · FFN 1,024",
                "256 → 1,024 → 1,000\n마스킹 위치의 Huber 손실 (delta = 1)"]
        geom_title="공간 그래프와 상대 기하 정보"
        geom_body="코어별 반경 층화 k200 / 500 um\n70차원 엣지 특성 → 헤드별 로짓 편향"
        mask="세포마다 0–1,000개 유전자 마스킹"
        residual="자기 루프 없음: 같은 세포의 정보는 잔차 연결로 전달"
        branch_title="탐색적 표현 분석"
        branch_body="최종 h4 → Leiden 군집화"
        future_title="향후 해석 분석"
        future_body="주의 가중치 프로파일 · Jacobian"
        footer="학습: 20개 코어 / 10개 기증자 그룹    검증: 8개 코어 / 4개 그룹    독립 테스트 없음"
        disclaimer="주의 가중치는 모델의 정보 전달 비율이며, 생물학적 인과 효과를 뜻하지 않음."
    else:
        header(s,"BAGM: masked expression learning in a tissue graph",
               "Current model: original Relative-Geometric QKV  |  Integrated SO_1 + SO_2 dataset",1)
        titles=["Masked cell inputs","Node encoder","4 graph blocks","Expression decoder"]
        bodies=["Visible-only Pearson features\nExplicit mask + 22 morphology fields",
                "Shared projection\nh⁰: 256-D", "Independent layer parameters\n8 heads · FFN width 1,024",
                "256 → 1,024 → 1,000\nMasked-entry Huber loss (δ = 1)"]
        geom_title="Spatial graph + relative geometry"
        geom_body="Per-core radial-stratified k200 / 500 μm\n70-D edge features → per-head logit bias"
        mask="Mask 0–1,000 genes independently per cell"
        residual="No self-loop edges; residual paths retain same-cell information"
        branch_title="Exploratory representation analysis"
        branch_body="Final h⁴ → Leiden clustering"
        future_title="Planned interpretation"
        future_body="Attention profiles · Jacobians"
        footer="Training: 20 cores / 10 donor groups    Validation: 8 cores / 4 groups    No independent test set"
        disclaimer="Attention describes computational routing; it does not by itself establish a biological or causal effect."
    s.rect(5.45,1.78,3.65,1.02,fill=MINT,stroke=TEAL)
    s.text(5.6,1.91,3.35,.28,geom_title,size=15,bold=True,align="center")
    s.text(5.58,2.27,3.4,.43,geom_body,size=12.3,align="center",color=GRAY)
    s.line(7.275,2.8,7.275,3.22,color=TEAL)
    specs=[(.5,2.65),(3.52,1.55),(5.45,3.65),(9.5,3.32)]
    for idx,(x,w) in enumerate(specs):
        s.rect(x,3.22,w,1.47,fill=PALE if idx!=3 else MINT,stroke=BLUE if idx!=3 else TEAL)
        s.text(x+.1,3.41,w-.2,.42,titles[idx],size=17 if idx!=1 else 14,bold=True,align="center")
        s.text(x+.1,4.0,w-.2,.6,bodies[idx],size=11 if idx==0 else (13 if idx==2 else 12),align="center",color=GRAY)
    for left,right in [(3.15,3.52),(5.07,5.45),(9.10,9.50)]:
        s.line(left,3.94,right,3.94)
    s.text(.55,2.63,4.55,.42,mask,size=12,color=GRAY)
    s.text(5.35,4.83,7.5,.3,residual,size=11,color=GRAY)
    s.line(7.275,4.69,7.275,5.45,color=TEAL,dash=True)
    s.rect(5.45,5.45,3.65,.9,fill=LIGHT,stroke=TEAL,dash=True)
    s.text(5.58,5.58,3.39,.3,branch_title,size=13.5,bold=True,align="center")
    s.text(5.58,5.96,3.39,.27,branch_body,size=12.5,align="center",color=GRAY)
    s.rect(9.5,5.45,3.32,.9,fill=LIGHT,stroke=RULE,dash=True)
    s.text(9.62,5.58,3.08,.3,future_title,size=14,bold=True,align="center")
    s.text(9.62,5.96,3.08,.27,future_body,size=12.5,align="center",color=GRAY)
    s.text(.55,5.48,4.48,.65,"Training objective" if not korean else "학습 목표",size=15,bold=True)
    training=("Predict hidden Pearson residual features\nfrom visible genes, morphology, and neighbors."
              if not korean else "관측 유전자, 형태 정보, 이웃 세포로\n숨겨진 Pearson 잔차 특성을 예측")
    s.text(.55,5.87,4.55,.5,training,size=13,color=GRAY)
    s.text(.55,6.57,12.25,.28,footer,size=11,color=GRAY)
    s.text(.55,6.89,12.25,.23,disclaimer,size=9.2,color=GRAY)
    s.notes(NOTES);s.end()


def equations(prs,pdf):
    s=Slide(prs,pdf)
    header(s,"Inside one Relative-Geometric QKV graph block",
           "Directed edge j → i: sender j supplies a value; receiver i selects among its incoming neighbors",2)
    cards=[(1.78,1.3,"1  Project the cell states",
            "uᵢ = LN(hᵢ);    qᵢₕ = WQ,ₕ uᵢ;    kⱼₕ = WK,ₕ uⱼ;    vⱼₕ = WV,ₕ uⱼ",
            "Each head has 32 dimensions; Q, K, and V use the normalized cell state."),
           (3.33,1.4,"2  Add geometry bias and aggregate neighbors",
            "sᵢⱼₕ = qᵢₕᵀ kⱼₕ / √32 + bᵢⱼₕ;    αᵢⱼₕ = softmaxⱼ∈N(i)(sᵢⱼₕ)",
            "zᵢ = concatₕ [Σⱼ∈N(i) αᵢⱼₕ vⱼₕ]"),
           (4.98,1.42,"3  Residual update and feed-forward network",
            "rᵢ = hᵢ + Dropout(WO zᵢ)",
            "hᵢ(next) = rᵢ + Dropout[FFN(LN(rᵢ))]     •     FFN: 256 → 1,024 → 256")]
    for y,h,title,eq,desc in cards:
        s.rect(.5,y,8.75,h,fill=PALE)
        s.text(.72,y+.13,8.31,.32,title,size=16,bold=True)
        s.text(.72,y+.55,8.31,.36,eq,size=17 if y<4 else 18)
        s.text(.72,y+.99,8.31,.31,desc,size=12.2 if y<3 else 15,color=GRAY)
    s.rect(9.58,1.78,3.25,2.95,fill=MINT,stroke=TEAL)
    s.text(9.78,1.97,2.85,.35,"Geometry branch",size=18,bold=True)
    s.text(9.78,2.46,2.85,.45,"ρᵢⱼ ∈ ℝ⁷⁰",size=23)
    s.text(9.78,3.04,2.85,.5,"70 → 128 → 8\nLinear · LN · GELU · Linear",size=14,color=GRAY)
    s.text(9.78,3.82,2.85,.58,"bᵢⱼₕ = fₕ(ρᵢⱼ)\nOne scalar bias per head",size=17)
    s.line(9.58,4.04,9.25,4.04,color=TEAL)
    s.rect(9.58,4.98,3.25,1.42,fill=LIGHT)
    s.text(9.78,5.14,2.85,.32,"4 independent graph blocks",size=14,bold=True)
    s.text(9.78,5.64,2.85,.7,"No implicit self edges\nGeometry enters scores only\nResidual/FFN dropout: 0.1",size=13.5,color=GRAY)
    s.text(.55,6.68,12.25,.3,"Current configuration: attention dropout = 0; residual and FFN dropout = 0.1. Geometry bias is zero at initialization.",size=11,color=GRAY)
    s.notes(NOTES+"\nEquation conventions: current configured attention dropout is 0.0, so the softmax weights enter value aggregation directly. Residual/FFN dropout is 0.1. RelativePositionBiasEncoder maps 70 to 128, LayerNorm and GELU, then 128 to 8. Its final projection is exactly zero-initialized, so geometry bias starts at zero. Values have no direct edge-geometry projection. The block is pre-LayerNorm and the residual FFN has GELU and dropout. The node encoder adds separate expression, mask, and morphology projections before LayerNorm, GELU, and dropout. The current 22 morphology features do not include additional missing-indicator channels.")
    s.end()


def paper_figure():
    w,h=6.8*72,1.55*72
    path=OUT/"BAGM_architecture_figure.pdf"
    c=canvas.Canvas(str(path),pagesize=(w,h))
    c.setTitle("BAGM original Relative-Geometric QKV architecture")
    def box(x,y,bw,bh,title,body,fill=PALE,stroke=BLUE,fs=8.4):
        c.setFillColorRGB(*rgb(fill));c.setStrokeColorRGB(*rgb(stroke));c.setLineWidth(.65)
        c.roundRect(x*72,y*72,bw*72,bh*72,4,fill=1,stroke=1)
        c.setFillColorRGB(*rgb(NAVY));c.setFont("FreeSansBold",9)
        c.drawCentredString((x+bw/2)*72,(y+bh-.18)*72,title)
        c.setFont("FreeSans",fs);c.setFillColorRGB(*rgb(GRAY))
        for j,line in enumerate(body.split("\n")):
            c.drawCentredString((x+bw/2)*72,(y+bh-.37)*72-j*9.8,line)
    def arrow(x1,y1,x2,y2,color=BLUE,dash=False):
        c.setStrokeColorRGB(*rgb(color));c.setFillColorRGB(*rgb(color));c.setLineWidth(.8)
        c.setDash(2,2) if dash else c.setDash()
        c.line(x1*72,y1*72,x2*72,y2*72);c.setDash()
        a=math.atan2(y2-y1,x2-x1);bx=x2*72;by=y2*72
        p=c.beginPath();p.moveTo(bx,by)
        p.lineTo(bx-4*math.cos(a)+1.7*math.sin(a),by-4*math.sin(a)-1.7*math.cos(a))
        p.lineTo(bx-4*math.cos(a)-1.7*math.sin(a),by-4*math.sin(a)+1.7*math.cos(a))
        p.close();c.drawPath(p,fill=1,stroke=0)
    box(.06,.40,1.56,.63,"Masked cell inputs","Visible-only Pearson + mask\n22 morphology fields",fs=8)
    box(1.88,.40,1.01,.63,"Node encoder","Shared projection\nh⁰ ∈ ℝ²⁵⁶",fs=8)
    box(3.15,.40,1.71,.63,"4 Relative-QKV blocks","8 heads; hidden width 256\nResidual + FFN (1,024)",fs=8)
    box(5.12,.40,1.62,.63,"Expression decoder","h⁴ → 1,024 → 1,000\nMasked Huber loss (δ = 1)",fill=MINT,stroke=TEAL,fs=8)
    for a,b in [(1.62,1.88),(2.89,3.15),(4.86,5.12)]:arrow(a,.715,b,.715)
    box(3.15,1.12,1.71,.40,"Spatial graph: k200 / 500 μm","70-D geometry → logit bias",fill=MINT,stroke=TEAL,fs=7.3)
    arrow(4.005,1.12,4.005,1.03,color=TEAL)
    c.setFillColorRGB(*rgb(GRAY));c.setFont("FreeSans",8)
    c.drawString(.08*72,1.20*72,"Mask 0–1,000 genes per cell")
    c.setFont("FreeSans",7.7)
    c.drawString(.08*72,.11*72,"Original Relative-Geometric QKV; same-cell information retained through residuals")
    arrow(4.98,.715,4.98,.18,color=TEAL,dash=True)
    c.setFont("FreeSans",7.7)
    c.drawString(5.12*72,.11*72,"h⁴ → Leiden (exploratory)")
    c.showPage();c.save()
    doc=pymupdf.open(path)
    doc[0].get_pixmap(matrix=pymupdf.Matrix(300/72,300/72),alpha=False).save(OUT/"BAGM_architecture_figure.png")
    doc.close()


def validate_pptx(path):
    ns={"p":"http://schemas.openxmlformats.org/presentationml/2006/main",
        "a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
    counts=[]
    with zipfile.ZipFile(path) as z:
        slides=sorted(n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        assert len(slides)==3
        for name in slides:
            root=ET.fromstring(z.read(name))
            assert not root.findall(".//p:pic",ns), "Raster content unexpectedly found"
            shapes=len(root.findall(".//p:sp",ns)); connectors=len(root.findall(".//p:cxnSp",ns))
            assert shapes>15 and connectors>=1
            counts.append((name,shapes,connectors))
        assert len([n for n in z.namelist() if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")])==3
    check=Presentation(path)
    for slide in check.slides:
        for shp in slide.shapes:
            assert shp.left>=0 and shp.top>=0
            assert shp.left+shp.width<=check.slide_width+100 and shp.top+shp.height<=check.slide_height+100
    return counts


def main():
    register_fonts()
    prs=Presentation();prs.slide_width=Inches(WIDTH);prs.slide_height=Inches(HEIGHT)
    prs.core_properties.title="BAGM original Relative-Geometric QKV architecture"
    prs.core_properties.subject="Editable English/Korean model diagram and exact block equations"
    prs.core_properties.author="BAGM draft"
    proof=canvas.Canvas(str(OUT/"BAGM_architecture_preview.pdf"),pagesize=(WIDTH*72,HEIGHT*72))
    proof.setTitle("BAGM architecture — vector layout proof")
    overview(prs,proof);equations(prs,proof);overview(prs,proof,korean=True)
    proof.save();pptx=OUT/"BAGM_architecture_editable.pptx";prs.save(pptx)
    paper_figure()
    counts=validate_pptx(pptx)
    doc=pymupdf.open(OUT/"BAGM_architecture_preview.pdf")
    for i,page in enumerate(doc):
        page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25),alpha=False).save(OUT/f"BAGM_architecture_slide_{i+1}_preview.png")
    doc.close()
    print("Generated editable PPTX (3 slides), vector proof, and manuscript PDF/PNG.")
    print("Validation: native shapes/connectors, notes, page count, slide bounds:",counts)


if __name__=="__main__":
    main()
