"""Build matched bilingual PDF/DOCX initial drafts from structured source text.

Run from BAGM root with PYTHONPATH=src and the isolated document environment.
Only this report directory is written. Scientific runs and data are untouched.
"""
from pathlib import Path
from xml.sax.saxutils import escape
import json
import re
import hashlib
import importlib.util
import sys

# Load the lightweight path module without importing the GPU training package.
_path_spec = importlib.util.spec_from_file_location('bagm_paper_paths', Path(__file__).resolve().parents[2] / 'src/spatial_benchmark/paths.py')
_path_module = importlib.util.module_from_spec(_path_spec)
sys.modules[_path_spec.name] = _path_module
_path_spec.loader.exec_module(_path_module)
current_paths = _path_module.current_paths
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Image
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz

OUT = current_paths().report_root / 'paper_draft_20261016'
DATA = json.loads((OUT / 'manuscript_content.json').read_text())
ASSETS = OUT / 'assets'
PREVIEW = OUT / 'previews'
PREVIEW.mkdir(exist_ok=True)
PAGE_W, PAGE_H = 595.276, 841.89
MARGIN_X, TOP, BOTTOM = 46, 43, 44
WIDTH = PAGE_W - 2*MARGIN_X
INK, TEAL, MUTED = '#172E40', '#176F79', '#536675'

for name, path in [
    ('Body', '/usr/share/fonts/truetype/freefont/FreeSerif.ttf'),
    ('BodyBold', '/usr/share/fonts/truetype/freefont/FreeSerifBold.ttf'),
    ('Sans', '/usr/share/fonts/truetype/freefont/FreeSans.ttf'),
    ('SansBold', '/usr/share/fonts/truetype/freefont/FreeSansBold.ttf'),
    ('Korean', ASSETS/'NanumGothic-Regular.ttf'),
    ('KoreanBold', ASSETS/'NanumGothic-Bold.ttf')
]:
    pdfmetrics.registerFont(TTFont(name, str(path)))

def rich(text, font):
    """Keep Hangul native and use an explicit fallback for uncovered symbols."""
    spans, pending, active = [], '', font
    cmap = pdfmetrics.getFont(font).face.charToGlyph
    for ch in text:
        target = font if ord(ch) in cmap or ch.isspace() else 'Body'
        if target != active:
            if pending:
                spans.append(f'<font name="{active}">{escape(pending)}</font>')
            pending, active = '', target
        pending += ch
    if pending:
        spans.append(f'<font name="{active}">{escape(pending)}</font>')
    return ''.join(spans).replace('\n', '<br/>')

def style(kind, lang, size):
    ko = lang == 'ko'
    font = 'Korean' if ko else 'Body'
    values = dict(fontName=font, fontSize=size, leading=size*(1.46 if ko else 1.19),
                  textColor=colors.HexColor(INK), spaceAfter=5.0, wordWrap='CJK' if ko else None)
    if kind == 'title':
        values.update(fontName='KoreanBold' if ko else 'SansBold', fontSize=18 if ko else 20,
                      leading=24, spaceAfter=6, textColor=colors.HexColor(INK))
    elif kind == 'subtitle':
        values.update(fontName='Korean' if ko else 'Sans', fontSize=8, leading=11,
                      spaceAfter=13, textColor=colors.HexColor(MUTED))
    elif kind == 'heading':
        values.update(fontName='KoreanBold' if ko else 'SansBold', fontSize=size+1.3,
                      leading=size+4, spaceBefore=7, spaceAfter=5, textColor=colors.HexColor(TEAL))
    elif kind == 'caption':
        values.update(fontSize=size-1.2, leading=(size-1.2)*1.3, spaceAfter=7,
                      textColor=colors.HexColor(MUTED))
    elif kind == 'equation':
        values.update(fontName='Body', fontSize=size+0.3, leading=size*1.55,
                      spaceBefore=1, spaceAfter=3, alignment=1, wordWrap=None)
    elif kind == 'ref':
        values.update(fontName='Body', fontSize=9.0, leading=10.8,
                      spaceAfter=4.3, wordWrap=None)
    return ParagraphStyle(kind, **values)

def placeholder_pdf():
    path = ASSETS/'results_placeholder.pdf'
    c = canvas.Canvas(str(path), pagesize=(WIDTH, 65))
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont('SansBold', 7.5)
    c.drawString(0, 56, 'PLANNED COMPARISON — PLACEHOLDER, NO DATA SHOWN')
    labels = ['Independent\nannotation', 'Input\nh0', 'Intermediate\nh1–h3', 'Final\nh4', 'PCA', 'scVI']
    gap, bw = 6, (WIDTH-5*6)/6
    for i,label in enumerate(labels):
        x = i*(bw+gap)
        c.setFillColor(colors.HexColor('#F5F8FA'))
        c.setStrokeColor(colors.HexColor('#B4C6CE'))
        c.setDash(2,2)
        c.roundRect(x, 4, bw, 43, 3, stroke=1, fill=1)
        c.setFillColor(colors.HexColor(INK))
        c.setFont('Sans', 8)
        lines=label.split('\n')
        for j,line in enumerate(lines):
            c.drawCentredString(x+bw/2, 29-j*11 if len(lines)>1 else 24, line)
    c.save()
    pdf=fitz.open(path)
    pdf[0].get_pixmap(matrix=fitz.Matrix(3,3)).save(ASSETS/'results_placeholder.png')

def image_spec(name):
    if name == 'architecture':
        path=OUT/'BAGM_architecture_figure.png'
        if not path.exists():
            candidates=list(OUT.glob('*architecture*png'))+list(ASSETS.glob('*architecture*png'))
            if not candidates:
                raise FileNotFoundError('Architecture figure is not yet available')
            path=candidates[0]
        iw, ih = ImageReader(str(path)).getSize()
        return path, WIDTH, min(WIDTH*ih/iw, 118)
    return ASSETS/'results_placeholder.png', WIDTH, 65

def items_for_page(lang, page_index, size):
    items=[]
    content=DATA[lang]
    if page_index==0:
        for kind in ['title','subtitle']:
            s=style(kind,lang,size)
            items.append((Paragraph(rich(content[kind],s.fontName),s),s.spaceBefore,s.spaceAfter))
    for kind,text in content['pages'][page_index]:
        if kind=='figure':
            p,w,h=image_spec(text)
            items.append((Image(str(p),width=w,height=h,kind='proportional'),4,3))
        elif kind=='references':
            s=style('ref',lang,size)
            for r in DATA['references']:
                v=f"[{r['n']}] {r['text']} "
                mark=rich(v,s.fontName)+f'<link href="{r["url"]}" color="#176F79">{escape(r["doi"])}</link>'
                items.append((Paragraph(mark,s),0,s.spaceAfter))
        else:
            s=style(kind,lang,size)
            items.append((Paragraph(rich(text,s.fontName),s),s.spaceBefore,s.spaceAfter))
    return items

def height(items):
    return sum(before+obj.wrap(WIDTH,2000)[1]+after for obj,before,after in items)

def build_pdf(lang):
    output=OUT/f'BAGM_draft_{"English" if lang=="en" else "Korean"}.pdf'
    c=canvas.Canvas(str(output),pagesize=(PAGE_W,PAGE_H))
    c.setTitle(DATA[lang]['title'])
    c.setAuthor('BAGM research draft')
    c.setSubject('Initial manuscript; current original Relative-QKV model; preliminary results')
    layouts=[]
    for p in range(3):
        size=10.5 if lang=='en' else 9.7
        while True:
            items=items_for_page(lang,p,size)
            used=height(items)
            if used<=PAGE_H-TOP-BOTTOM-10:
                break
            size-=0.1
            if size<8.6:
                raise RuntimeError(f'{lang} page {p+1} cannot fit: {used}')
        y=PAGE_H-TOP
        for obj,before,after in items:
            y-=before
            _,h=obj.wrap(WIDTH,2000)
            obj.drawOn(c,MARGIN_X,y-h)
            y-=h+after
        c.setStrokeColor(colors.HexColor('#CDDAE0'))
        c.setLineWidth(.5)
        c.line(MARGIN_X,31,PAGE_W-MARGIN_X,31)
        c.setFont('Sans',7.2)
        c.setFillColor(colors.HexColor(MUTED))
        c.drawString(MARGIN_X,20,'BAGM  |  Initial draft  |  16 October 2026')
        c.drawRightString(PAGE_W-MARGIN_X,20,f'{p+1} / 3')
        layouts.append({'page':p+1,'body_pt':round(size,2),'used_pt':round(used,2),'bottom_y':round(y,2)})
        c.showPage()
    c.save()
    doc=fitz.open(output)
    for p in range(len(doc)):
        doc[p].get_pixmap(matrix=fitz.Matrix(1.4,1.4)).save(PREVIEW/f'{lang}_page{p+1}.png')
    return layouts

def set_font(run, name, size, bold=False):
    run.font.name=name
    run.font.size=Pt(size)
    run.bold=bold
    run.font.color.rgb=RGBColor.from_string(INK[1:])
    rpr=run._element.get_or_add_rPr()
    fonts=rpr.rFonts
    if fonts is not None:
        fonts.set(qn('w:eastAsia'),'NanumGothic')

def add_hyperlink(paragraph,label,url):
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),TEAL[1:]); rp.append(color)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'17'); rp.append(size)
    r.append(rp); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); paragraph._p.append(h)

def build_docx(lang,layouts):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Pt(PAGE_W); sec.page_height=Pt(PAGE_H)
    sec.top_margin=Pt(TOP); sec.bottom_margin=Pt(BOTTOM)
    sec.left_margin=Pt(MARGIN_X); sec.right_margin=Pt(MARGIN_X)
    sec.footer_distance=Pt(18)
    ko=lang=='ko'; font='NanumGothic' if ko else 'Times New Roman'
    normal=doc.styles['Normal']
    normal.font.name=font; normal.font.size=Pt(9.1 if ko else 10)
    normal.paragraph_format.space_after=Pt(4)
    doc.core_properties.title=DATA[lang]['title']
    doc.core_properties.subject='Initial draft with provisional observations and figure placeholder'
    doc.core_properties.author=''
    for pidx,page in enumerate(DATA[lang]['pages']):
        if pidx:
            doc.add_page_break()
        size=layouts[pidx]['body_pt']
        if not pidx:
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.line_spacing=Pt(23)
            set_font(p.add_run(DATA[lang]['title']),font,18 if ko else 20,True)
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(10)
            set_font(p.add_run(DATA[lang]['subtitle']),font,8)
        for kind,text in page:
            if kind=='figure':
                path,w,h=image_spec(text)
                p=doc.add_paragraph()
                p.paragraph_format.space_after=Pt(2)
                p.add_run().add_picture(str(path),width=Pt(w),height=Pt(h))
                continue
            if kind=='references':
                for ref in DATA['references']:
                    p=doc.add_paragraph()
                    p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.line_spacing=Pt(10)
                    set_font(p.add_run(f"[{ref['n']}] {ref['text']} "),'Times New Roman',8.5)
                    add_hyperlink(p,ref['doi'],ref['url'])
                continue
            p=doc.add_paragraph()
            s=style(kind,lang,size)
            p.paragraph_format.space_before=Pt(s.spaceBefore)
            p.paragraph_format.space_after=Pt(max(s.spaceAfter-1,0))
            p.paragraph_format.line_spacing=Pt(s.leading)
            p.paragraph_format.widow_control=True
            if kind=='heading':
                p.paragraph_format.keep_with_next=True
            if kind=='equation':
                p.alignment=1
            set_font(p.add_run(text),'Cambria Math' if kind=='equation' else font,s.fontSize,kind=='heading')
        footer=sec.footer.paragraphs[0]
    footer.text='BAGM | Initial draft | 16 October 2026'
    footer.runs[0].font.size=Pt(7)
    path=OUT/f'BAGM_draft_{"Korean" if ko else "English"}.docx'
    doc.save(path)

def verify():
    audit={}
    for lang,label in [('en','English'),('ko','Korean')]:
        path=OUT/f'BAGM_draft_{label}.pdf'
        d=fitz.open(path)
        text='\n'.join(p.get_text() for p in d)
        assert len(d)==3,(path,len(d))
        assert '407,659' in text and 'Relative' in text
        assert all(f'[{i}]' in text for i in range(1,11))
        assert len([x for p in d for x in p.get_links()])>=10
        if lang=='ko':
            assert '공간' in text and '참고문헌' in text
        outside=[]
        for idx,p in enumerate(d):
            for b in p.get_text('blocks'):
                if b[0]<0 or b[1]<0 or b[2]>PAGE_W+.1 or b[3]>PAGE_H+.1:
                    outside.append([idx,b[:4]])
        assert not outside,outside
        docx=Document(OUT/f'BAGM_draft_{label}.docx')
        assert len(docx.inline_shapes)==2
        audit[label]={'pages':len(d),'text_characters':len(text),'links':len([x for p in d for x in p.get_links()]),'page_bounds_ok':True,'docx_figures':len(docx.inline_shapes)}
    combined=fitz.open()
    for lang in ['English','Korean']:
        combined.insert_pdf(fitz.open(OUT/f'BAGM_draft_{lang}.pdf'))
    combined.save(OUT/'BAGM_draft_bilingual.pdf')
    return audit

if __name__=='__main__':
    placeholder_pdf()
    layouts={}
    for language in ['en','ko']:
        layouts[language]=build_pdf(language)
        build_docx(language,layouts[language])
    result={'layout':layouts,'validation':verify()}
    (OUT/'document_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False,indent=2))
